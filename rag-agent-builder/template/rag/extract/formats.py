"""PDF 이외 형식 추출 -> PDF와 같은 페이지 dict (page_no, page_type, char_count, blocks)

쪽이 없는 형식은 page_no를 '위치 번호'로 쓴다 (unit에 기록): 근거 위치·골든셋 쪽 번호가 이 번호를 따른다.
| 형식 | 위치 단위 | 블록 |
|---|---|---|
| .docx | 제목(Heading 1·2)마다 또는 SECTION_CHARS마다 나눈 section | text, table |
| .pptx | slide | text, table |
| .xlsx .csv | sheet (csv는 1) | table (첫 행 = 헤더) |
| .html .htm | h1~h3 제목마다 section | text, table |
| .md .txt | # 제목마다 또는 SECTION_CHARS마다 section | markdown (표는 청킹 때 행 단위) |
| 이미지 | 파일 하나 = page 1 | image (needs_vision -> ocr.py) |
| .hwp .hwpx .doc .ppt .xls .rtf .odt | LibreOffice로 PDF 변환 후 PDF 규칙 (없으면 .hwpx만 본문 텍스트) |
"""

import csv
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from rag import config
from rag.common import is_real_table, markdown_table

SECTION_CHARS = 3000  # 제목이 없는 긴 문서는 이 길이마다 위치 번호를 나눈다
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
CONVERT_SUFFIXES = {".hwp", ".hwpx", ".doc", ".ppt", ".xls", ".rtf", ".odt", ".odp", ".ods"}


def page(page_no, blocks, unit, image=False):
    chars = sum(len(b.get("text", "")) for b in blocks)
    page_type = "scan" if image else "text" if chars else "empty"
    for i, b in enumerate(blocks):
        b["block_id"] = f"page-{page_no}-block-{i}"
    return {"page_no": page_no, "page_type": page_type, "unit": unit, "char_count": chars, "blocks": blocks}


def text_block(text):
    return {"content_type": "text", "text": text}


def table_block(rows, table_id):
    """2차원 리스트 -> 표 블록. 1행/1열은 본문 텍스트로"""
    rows = [[" ".join(str(c or "").split()) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not is_real_table(rows):
        return text_block("\n".join(" ".join(r) for r in rows)) if rows else None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return {"content_type": "table", "table_id": table_id, "headers": rows[0], "rows": rows[1:], "footnotes": "",
            "text": markdown_table(rows[0], rows[1:])}


def sections(items, unit):
    """[(블록, 제목 여부)] -> 제목이나 SECTION_CHARS 기준으로 나눈 페이지 목록"""
    pages, cur, size, has_body = [], [], 0, False
    for block, heading in items:
        if has_body and (heading or size >= SECTION_CHARS):  # 제목만 연달아 있으면 나누지 않는다
            pages.append(cur)
            cur, size, has_body = [], 0, False
        cur.append(block)
        size += len(block.get("text", ""))
        has_body = has_body or not heading
    pages += [cur] if cur else []
    return [page(i, merge_text(bs), unit) for i, bs in enumerate(pages, start=1)]


def merge_text(blocks):
    """이어진 text 블록은 하나로"""
    out = []
    for b in blocks:
        if out and b["content_type"] == out[-1]["content_type"] == "text":
            out[-1]["text"] += "\n" + b["text"]
        else:
            out.append(b)
    return out


# ---------------- 형식별 ----------------

def docx_pages(path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc, items = Document(path), []
    for i, el in enumerate(doc.element.body.iterchildren()):
        if el.tag.endswith("}p"):
            p = Paragraph(el, doc)
            if p.text.strip():
                style = (p.style.name or "") if p.style is not None else ""
                items.append((text_block(p.text.strip()), bool(re.match(r"(Heading|제목) ?[12]$", style))))
        elif el.tag.endswith("}tbl"):
            t = table_block([[c.text for c in row.cells] for row in Table(el, doc).rows], f"table-{i}")
            items += [(t, False)] if t else []
    return sections(items, "section")


def pptx_pages(path):
    from pptx import Presentation
    pages = []
    for n, slide in enumerate(Presentation(path).slides, start=1):
        blocks = []
        for s_idx, shape in enumerate(sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0))):
            if shape.has_text_frame and shape.text_frame.text.strip():
                blocks.append(text_block(shape.text_frame.text.strip()))
            elif getattr(shape, "has_table", False) and shape.has_table:
                t = table_block([[c.text for c in r.cells] for r in shape.table.rows], f"slide-{n}-table-{s_idx}")
                blocks += [t] if t else []
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            blocks.append(text_block("[발표자 노트]\n" + slide.notes_slide.notes_text_frame.text.strip()))
        pages.append(page(n, merge_text(blocks), "slide"))
    return pages


def sheet_pages(path):
    if path.suffix.lower() == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            sheets = [("csv", list(csv.reader(f)))]
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    pages = []
    for n, (name, rows) in enumerate(sheets, start=1):
        t = table_block(rows, f"sheet-{n}")
        blocks = [text_block(f"[시트] {name}")] + ([t] if t else [])
        pages.append(page(n, blocks, "sheet"))
    return pages


def html_pages(path):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer", "aside", "form"]):
        tag.decompose()
    items = []
    for i, el in enumerate(soup.find_all(["h1", "h2", "h3", "p", "li", "pre", "table", "dt", "dd"])):
        if el.find_parent(["table", "li", "p", "pre"]):
            continue  # 표·목록·문단 안의 요소는 바깥 요소에서 한 번만 읽는다
        if el.name == "table":
            t = table_block([[c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])]
                             for tr in el.find_all("tr")], f"table-{i}")
            items += [(t, False)] if t else []
        elif text := el.get_text(" ", strip=True):
            items.append((text_block(text), el.name in ("h1", "h2", "h3")))
    return sections(items, "section")


def text_pages(path):
    text = path.read_text(encoding="utf-8", errors="ignore")
    parts = re.split(r"\n(?=#{1,3} )", text) if path.suffix.lower() == ".md" else re.split(r"\n\s*\n", text)
    items = [({"content_type": "markdown", "text": p.strip()}, p.lstrip().startswith("#")) for p in parts if p.strip()]
    return sections(items, "section")


def image_pages(path):
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    block = {"content_type": "image", "image_id": "page-1-image-full", "bbox": [0, 0, w, h], "area_ratio": 1.0,
             "text": "", "needs_vision": True}
    return [page(1, [block], "image", image=True)]


def hwpx_text_pages(path):
    """LibreOffice가 없을 때 .hwpx 본문 글자만 (표 구조·이미지는 잃는다)"""
    with zipfile.ZipFile(path) as z:
        xml = "".join(z.read(n).decode("utf-8") for n in sorted(z.namelist()) if re.match(r"Contents/section\d+\.xml", n))
    paras = [re.sub(r"<[^>]+>", "", p) for p in re.findall(r"<hp:p\b.*?</hp:p>", xml, flags=re.S)]
    return sections([(text_block(p.strip()), False) for p in paras if p.strip()], "section")


# ---------------- 변환 ----------------

def soffice():
    found = shutil.which("soffice") or shutil.which("libreoffice")
    default = Path("C:/Program Files/LibreOffice/program/soffice.exe")
    return found or (str(default) if default.exists() else None)


def convert_to_pdf(path):
    """LibreOffice로 PDF 변환 (output/converted/에 캐시). 실패하면 None"""
    exe = soffice()
    if not exe:
        return None
    out_dir = config.OUT_DIR / "converted" / path.parent.relative_to(config.DATA_DIR)
    pdf = out_dir / f"{path.stem}.pdf"
    if not pdf.exists():
        out_dir.mkdir(parents=True, exist_ok=True)
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(path)],
                       check=False, capture_output=True, timeout=300)
    return pdf if pdf.exists() else None


EXTRACTORS = {".docx": docx_pages, ".pptx": pptx_pages, ".xlsx": sheet_pages, ".xlsm": sheet_pages,
              ".csv": sheet_pages, ".html": html_pages, ".htm": html_pages, ".md": text_pages, ".txt": text_pages,
              **{s: image_pages for s in IMAGE_SUFFIXES}}
